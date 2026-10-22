#!/usr/bin/env python3
"""Genera dist/Observatorio_Violencia_Armada_Colombia_2026.xlsx con indicadores calculados por fórmulas."""
import json, subprocess, sys, datetime as dt
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
subprocess.run([sys.executable, str(ROOT / "src" / "validate.py")], check=True, stdout=subprocess.DEVNULL)
D = json.loads((ROOT / "data" / "data.json").read_text(encoding="utf-8"))
EV = D["events"]
N = len(EV)
LAST = N + 1  # última fila de datos en 'Incidentes'

AR = "Arial"
HF = PatternFill("solid", fgColor="0D1F3C")
HFONT = Font(name=AR, bold=True, color="FFFFFF", size=10)
BASE = Font(name=AR, size=10)
BOLD = Font(name=AR, size=10, bold=True)
LINK = Font(name=AR, size=10, color="0053BB", underline="single")
TITLE = Font(name=AR, size=16, bold=True, color="0D1F3C")
SUB = Font(name=AR, size=11, color="46556D")
thin = Side(style="thin", color="D6DDE8")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.alignment, c.border = HFONT, HF, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def put(ws, r, c, v, font=BASE, fmt=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = WRAP
    return cell


# ---------------------------------------------------------------- Portada
ws = wb.active
ws.title = "Portada"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 80
logo = ROOT / "assets" / "logo_geseg_azul.png"
if logo.exists():
    img = XImage(str(logo))
    img.width, img.height = 220, int(220 * img.height / img.width)
    ws.add_image(img, "B2")
ws["B8"] = "Observatorio territorial de terrorismo y violencia armada en Colombia"
ws["B8"].font = TITLE
ws["B9"] = "Base de datos depurada · Producto D"
ws["B9"].font = SUB
rows = [("Periodo", "Desde el 7 de agosto de 2026 (00:00, UTC−5)"), ("Fecha de corte", D["updated"]),
        ("Registros en la base", f"=COUNTA(Incidentes!A2:A{LAST})"),
        ("Hojas", "Portada · Incidentes · Indicadores · Fuentes · Excluidos y pendientes · Metodología · Diccionario"),
        ("Elaboración", "GESEG – Gestión de la Seguridad S.A.S. · Del dato a la decisión."),
        ("Nota", "Instrumento técnico con fuentes abiertas. No es producto de inteligencia ni de contrainteligencia. Los conteos por actor no equivalen a responsabilidad judicial. Revisión independiente pendiente de asignación.")]
for i, (k, v) in enumerate(rows, 11):
    ws.cell(row=i, column=2, value=k).font = BOLD
    c = ws.cell(row=i, column=3, value=v)
    c.font = BASE
    c.alignment = WRAP
ws.row_dimensions[16].height = 48

# ---------------------------------------------------------------- Incidentes
wi = wb.create_sheet("Incidentes")
COLS = [  # (encabezado, clave, ancho)
    ("Id", "id", 13), ("Fecha", "fecha", 12), ("Hora", "hora", 9), ("Departamento", "depto", 18), ("Municipio", "mpio", 18),
    ("Lugar", "lugar", 34), ("Lat", "lat", 8), ("Lon", "lon", 9), ("Precisión coordenada", "prec_coord", 20), ("Hecho", "tipo", 44),
    ("Modalidad", "mod", 24), ("Objetivo", "objetivo", 26), ("Clase (analítica)", "clase", 34), ("Presunto autor", "autor", 40),
    ("Grupo", "grupo", 22), ("Atribución", "atrib", 10),
    ("Muertos civiles", "m_civ", 10), ("Muertos FP", "m_fp", 10), ("Heridos civiles", "h_civ", 10), ("Heridos FP", "h_fp", 10),
    ("Total muertos", None, 10), ("Total heridos", None, 10), ("Nota de víctimas", "h_nota", 36), ("Evidencia", "evidencia", 40),
    ("Daños", "danos", 36), ("Económico", "eco", 16), ("Desplazamiento", "desplaz", 16), ("Servicios", "servicios", 26),
    ("Estado", "estado", 16), ("Fuente 1", "f1", 30), ("Fuente 2", "f2", 30), ("Fuente 3", "f3", 30), ("N.º fuentes", None, 9),
    ("Última verificación", "actualizacion", 14)]
header(wi, 1, [c[0] for c in COLS], [c[2] for c in COLS])
wi.row_dimensions[1].height = 32
for r, e in enumerate(EV, 2):
    for ci, (h, k, w) in enumerate(COLS, 1):
        if h == "Total muertos":
            put(wi, r, ci, f"=Q{r}+R{r}")
        elif h == "Total heridos":
            put(wi, r, ci, f"=S{r}+T{r}")
        elif h == "N.º fuentes":
            put(wi, r, ci, f'=COUNTIF(AD{r}:AF{r},"?*")')
        elif k in ("f1", "f2", "f3"):
            i = int(k[1]) - 1
            u = e["fuentes"][i] if i < len(e["fuentes"]) else None
            c = put(wi, r, ci, u)
            if u:
                c.hyperlink, c.font = u, LINK
        elif k == "fecha":
            put(wi, r, ci, dt.date.fromisoformat(e[k]), fmt="yyyy-mm-dd")
        elif k in ("lat", "lon"):
            put(wi, r, ci, e[k], fmt="0.00")
        else:
            put(wi, r, ci, e[k], wrap=isinstance(e[k], str) and len(e[k]) > 30)
wi.freeze_panes = "F2"
tb = Table(displayName="Incidentes", ref=f"A1:{L(len(COLS))}{LAST}")
tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
wi.add_table(tb)

# ---------------------------------------------------------------- Indicadores
wx = wb.create_sheet("Indicadores")
R = lambda col: f"Incidentes!${col}$2:${col}${LAST}"
wx.column_dimensions["A"].width = 46
for c in "BCDEFGHIJ":
    wx.column_dimensions[c].width = 15
wx["A1"] = "Indicadores (todas las cifras se calculan con fórmulas sobre la hoja Incidentes)"
wx["A1"].font = TITLE
header(wx, 3, ["Indicador", "Valor"])
ind = [
    ("Incidentes", f"=COUNTA({R('A')})", "0"),
    ("Fallecidos", f"=SUM({R('U')})", "0"),
    ("  de ellos, Fuerza Pública", f"=SUM({R('R')})", "0"),
    ("  de ellos, civiles", f"=SUM({R('Q')})", "0"),
    ("Heridos", f"=SUM({R('V')})", "0"),
    ("  de ellos, Fuerza Pública", f"=SUM({R('T')})", "0"),
    ("  de ellos, civiles", f"=SUM({R('S')})", "0"),
    ("Civiles afectados (fallecidos + heridos)", "=B7+B10", "0"),
    ("Departamentos con incidentes", f"=SUMPRODUCT(1/COUNTIF({R('D')},{R('D')}))", "0"),
    ("Municipios con incidentes", f"=SUMPRODUCT(1/COUNTIFS({R('D')},{R('D')},{R('E')},{R('E')}))", "0"),
    ("Atribuidos al ELN (presunto)", f'=COUNTIF({R("O")},"ELN")', "0"),
    ("Atribuidos a disidencias de las FARC (presunto)", f'=COUNTIF({R("O")},"Disidencias FARC")', "0"),
    ("Otros actores", f'=B4-B14-B15-COUNTIF({R("O")},"No determinado")', "0"),
    ("Autoría desconocida (atribución E)", f'=COUNTIF({R("P")},"E")', "0"),
    ("% sin autor determinado", "=B17/B4", "0.0%"),
    ("Registros con fuente única", f"=COUNTIF({R('AG')},1)", "0"),
    ("% con información incompleta (fuente única, preliminar o en investigación)",
     f'=SUMPRODUCT(--((({R("AG")}=1)+ISNUMBER(SEARCH("Preliminar",{R("AC")}))+ISNUMBER(SEARCH("investigación",{R("AC")})))>0))/B4', "0.0%"),
]
for i, (k, f, fmt) in enumerate(ind, 4):
    put(wx, i, 1, k)
    put(wx, i, 2, f, fmt=fmt)
r0 = 4 + len(ind) + 2
GRS = ["ELN", "Disidencias FARC", "Clan del Golfo / EGC", "ACSN", "Autor individual (capturado)", "No determinado"]
header(wx, r0, ["Por grupo", "Incidentes", "Fallecidos", "Heridos"])
for i, g in enumerate(GRS, r0 + 1):
    put(wx, i, 1, g)
    put(wx, i, 2, f'=COUNTIF({R("O")},A{i})')
    put(wx, i, 3, f'=SUMIF({R("O")},A{i},{R("U")})')
    put(wx, i, 4, f'=SUMIF({R("O")},A{i},{R("V")})')
r1 = r0 + len(GRS) + 3
header(wx, r1, ["Por tipo de atribución", "Incidentes"])
ATRL = {"A": "A · Judicial", "B": "B · Oficial", "C": "C · Reivindicada", "D": "D · Hipótesis periodística o de terceros", "E": "E · Desconocida"}
for i, (k, v) in enumerate(ATRL.items(), r1 + 1):
    put(wx, i, 1, v)
    put(wx, i, 2, f'=COUNTIF({R("P")},"{k}")')
r2 = r1 + len(ATRL) + 3
deps = sorted({e["depto"] for e in EV})
header(wx, r2, ["Por departamento", "Incidentes", "Fallecidos", "Heridos", "Sin autor (E)", "% sin autor"])
for i, d in enumerate(deps, r2 + 1):
    put(wx, i, 1, d)
    put(wx, i, 2, f'=COUNTIF({R("D")},A{i})')
    put(wx, i, 3, f'=SUMIF({R("D")},A{i},{R("U")})')
    put(wx, i, 4, f'=SUMIF({R("D")},A{i},{R("V")})')
    put(wx, i, 5, f'=COUNTIFS({R("D")},A{i},{R("P")},"E")')
    put(wx, i, 6, f"=E{i}/B{i}", fmt="0%")
r3 = r2 + len(deps) + 3
header(wx, r3, ["Por mes", "Desde", "Hasta (excl.)", "Incidentes", "Fallecidos", "Heridos"])
months = sorted({e["fecha"][:7] for e in EV})
for i, m in enumerate(months, r3 + 1):
    y, mo = map(int, m.split("-"))
    nxt = dt.date(y + (mo == 12), mo % 12 + 1, 1)
    put(wx, i, 1, ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"][mo] + (" (desde el 7)" if m == "2026-08" else ""))
    put(wx, i, 2, dt.date(2026, 8, 7) if m == "2026-08" else dt.date(y, mo, 1), fmt="yyyy-mm-dd")
    put(wx, i, 3, nxt, fmt="yyyy-mm-dd")
    put(wx, i, 4, f'=COUNTIFS({R("B")},">="&B{i},{R("B")},"<"&C{i})')
    put(wx, i, 5, f'=SUMIFS({R("U")},{R("B")},">="&B{i},{R("B")},"<"&C{i})')
    put(wx, i, 6, f'=SUMIFS({R("V")},{R("B")},">="&B{i},{R("B")},"<"&C{i})')
wx.freeze_panes = "B4"

# ---------------------------------------------------------------- Fuentes
wf = wb.create_sheet("Fuentes")
header(wf, 1, ["Id", "URL", "Dominio"], [14, 100, 28])
r = 2
for e in EV:
    for u in e["fuentes"]:
        put(wf, r, 1, e["id"])
        c = put(wf, r, 2, u)
        c.hyperlink, c.font = u, LINK
        put(wf, r, 3, urlparse(u).netloc.replace("www.", ""))
        r += 1
r += 1
wf.cell(row=r, column=1, value="Contexto agregado (no se suma a la base)").font = BOLD
r += 1
for i, h in enumerate(["Fuente", "Cifra agregada descrita", "URL"], 1):
    c = wf.cell(row=r, column=i, value=h)
    c.font, c.fill = HFONT, HF
r += 1
wf.column_dimensions["B"].width = 100
for f, cifra, u in D["context"]:
    put(wf, r, 1, f, wrap=True)
    put(wf, r, 2, cifra, wrap=True)
    c = put(wf, r, 3, u)
    c.hyperlink, c.font = u, LINK
    r += 1
wf.column_dimensions["A"].width = 36
wf.freeze_panes = "A2"

# ---------------------------------------------------------------- Excluidos
we = wb.create_sheet("Excluidos y pendientes")
header(we, 1, ["Fecha", "Departamento", "Municipio", "Hecho reportado", "Justificación"], [12, 20, 22, 60, 70])
for r, x in enumerate(D["excluded"], 2):
    for c, v in enumerate(x, 1):
        put(we, r, c, v, wrap=True)
we.freeze_panes = "A2"

# ---------------------------------------------------------------- Metodología
wm = wb.create_sheet("Metodología")
wm.column_dimensions["A"].width = 28
wm.column_dimensions["B"].width = 120
header(wm, 1, ["Tema", "Criterio"])
MET = [
    ("1. Alcance temporal y territorial", "Hechos ocurridos en Colombia desde el 7 de agosto de 2026 (00:00, UTC−5) hasta la fecha de corte."),
    ("2. Alcance temático", "Presuntos actos terroristas, atentados, ataques armados y otras acciones violentas relevantes. No se equipara el terrorismo con cualquier violencia."),
    ("3. Fuentes", "Fuentes abiertas: entidades oficiales, agencias y medios nacionales y regionales. Se leyó cada nota citada; las fuentes que no pudieron consultarse no se citan como leídas."),
    ("4. Inclusión", "Un registro por evento verificable con fecha, lugar y modalidad. Los registros con fuente única o estado preliminar se marcan como información incompleta."),
    ("5. Clasificación", "La clase analítica es provisional y no es calificación jurídica ni denominación oficial. Se distingue entre calificación jurídica, denominación oficial y clasificación analítica."),
    ("6. Atribución A–E", "A judicial · B oficial · C reivindicada · D hipótesis periodística o de terceros · E desconocida. No se infiere autoría por presencia territorial, modalidad, antecedentes o coincidencia geográfica."),
    ("7. Víctimas", "Cifra oficial más reciente; No se publican nombres de víctimas."),
    ("8. Coordenadas", "Centroide municipal aproximado (WGS84), nunca una ubicación exacta."),
    ("9. Daño económico", "No se asignan valores económicos sin fuente con metodología. Una recompensa no es un daño."),
    ("10. Deduplicación", "Un registro por evento; varios medios sobre el mismo hecho se consolidan en un solo id. Se excluyen combates, atentados frustrados, hechos fuera del periodo o del territorio y fechas no verificables."),
    ("11. Vacíos y privacidad", "La base es una muestra, no el universo. Las cifras agregadas oficiales van aparte y no se suman. No se publica información táctica que facilite ataques."),
]
for r, (a, b) in enumerate(MET, 2):
    put(wm, r, 1, a, BOLD, wrap=True)
    put(wm, r, 2, b, wrap=True)

# ---------------------------------------------------------------- Diccionario
wd = wb.create_sheet("Diccionario")
header(wd, 1, ["Campo", "Tipo", "Descripción", "Valores o regla"], [18, 10, 70, 70])
DIC = [
    ("id", "texto", "Identificador consecutivo, nunca se reutiliza", "CO-2026-NNN"),
    ("fecha", "fecha", "Fecha de ocurrencia (no de publicación)", "AAAA-MM-DD, desde 2026-08-07 hasta el corte"),
    ("hora", "texto", "Hora o franja", "HH:MM, franja o 'Sin dato'"),
    ("depto, mpio, lugar", "texto", "Ubicación; lugar es vereda, barrio o sector si está documentado", ""),
    ("lat, lon", "número", "Centroide municipal aproximado (WGS84)", "Lat −4,3 a 13,5; lon −79,1 a −66,8"),
    ("prec_coord", "texto", "Precisión de la coordenada", "Municipal aproximada · Departamental aproximada"),
    ("tipo", "texto", "Modalidad descrita por las fuentes", ""),
    ("mod", "texto", "Categoría controlada", "Dron con explosivos · Vehículo con explosivos · Explosivo / IED · Granada / artefacto lanzado · Paro armado · Homicidio selectivo · Ataque armado"),
    ("objetivo", "texto", "Población, Fuerza Pública, infraestructura u otro", ""),
    ("clase", "texto", "Clasificación analítica provisional", ""),
    ("autor", "texto", "Actor mencionado por las fuentes", ""),
    ("grupo", "texto", "Agregado para los indicadores", "ELN · Disidencias FARC · Clan del Golfo / EGC · ACSN · Autor individual (capturado) · No determinado"),
    ("atrib", "texto", "Tipo de atribución", "A judicial · B oficial · C reivindicada · D hipótesis · E desconocida"),
    ("evidencia", "texto", "Fundamento público de la atribución", ""),
    ("m_civ, m_fp, h_civ, h_fp", "entero", "Muertos y heridos, civiles y Fuerza Pública; cifra oficial más reciente", "≥ 0"),
    ("h_nota", "texto", "Precisiones sobre las víctimas", ""),
    ("danos, eco, desplaz, servicios", "texto", "Daños verificables; eco = 'No cuantificado' si no hay fuente", ""),
    ("estado", "texto", "Estado del registro", "Confirmado · Preliminar · En investigación"),
    ("fuentes", "lista", "1 a 3 URL verificables", "https://"),
    ("actualizacion", "fecha", "Fecha de la última verificación", ""),
    ("Total muertos / heridos / N.º fuentes", "fórmula", "Calculados en la hoja Incidentes", "Fórmulas: Q+R · S+T · COUNTIF(AD:AF,\"?*\")"),
]
for r, row in enumerate(DIC, 2):
    for c, v in enumerate(row, 1):
        put(wd, r, c, v, wrap=True)

out = ROOT / "dist" / "Observatorio_Violencia_Armada_Colombia_2026.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print("OK ->", out, f"({len(EV)} registros, {len(wb.sheetnames)} hojas)")
